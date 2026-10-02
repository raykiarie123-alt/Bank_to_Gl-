#!/usr/bin/env python3
"""
bank_to_gl.py  -  Bank statement  ->  draft General Ledger journal (D365 template format)

What it does
  1. Reads the bank statement (Excel), cleans dates / amounts, skips opening-balance rows
  2. Codes each line using GL_Rules.xlsx (keyword rules) - reversals are detected automatically
  3. Builds balanced 2-line journals in your GL template layout (import file)
  4. Runs checks (balanced, totals agree with the statement, duplicates, unclassified)
  5. Writes a review workbook: Summary, Needs Review, All Lines and (optionally) a test against
     the GL entries you already posted for the same month

Nothing is ever posted. The output is a DRAFT file to review and import as an unposted journal.

Example (January test):
  python bank_to_gl.py --statement Automation.xlsx --sheet "Jan 2026" --rules GL_Rules.xlsx \
      --start-doc 300 --compare-gl "General Ledger Entries Jan" --out-dir out
"""
import argparse, re, datetime as dt, calendar
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

GL_COLUMNS = ['Document No.', 'External Document No.', 'G/L Account No.', 'Description', 'cost Center Code',
              'Cost Center Name', 'Project Code', 'Project Name', 'Donor code ', 'Donor Name', 'Column1', 'Column2',
              'Posting Date', 'Document Type', 'Debit Amount', 'Credit Amount', 'Bal. Account Type',
              'Bal. Account No.', 'Reversed', 'Reversed by Entry No.', 'Entry No.', 'Column4', 'Column3']

# Statement layout (0-based column positions). Matches your January sheet.
C_DATE, C_REF, C_NARR, C_WD, C_DEP = 0, 2, 3, 5, 6

CLEARING_ACCOUNT_SETTING = {'PAYE': 'PAYE Clearing Account', 'NSSF': 'NSSF Clearing Account',
                             'SHIF': 'SHIF Clearing Account', 'HOUSING': 'Housing Levy Clearing Account'}

YELLOW = PatternFill('solid', fgColor='FFF2CC'); ORANGE = PatternFill('solid', fgColor='F8CBAD')
GREEN = PatternFill('solid', fgColor='E2EFDA'); NAVY = PatternFill('solid', fgColor='1F3864')
WHITE_B = Font(bold=True, color='FFFFFF')


# ---------------------------------------------------------------- helpers
def num(v):
    if v is None or (isinstance(v, float) and np.isnan(v)): return 0.0
    s = str(v).replace(',', '').strip()
    if s in ('', '-'): return 0.0
    try: return float(s)
    except ValueError: return 0.0

def to_date(v):
    if isinstance(v, (dt.datetime, pd.Timestamp)): return pd.Timestamp(v).normalize()
    if isinstance(v, str):
        for f in ('%d-%m-%Y', '%Y-%m-%d', '%d/%m/%Y'):
            try: return pd.Timestamp(dt.datetime.strptime(v.strip(), f))
            except ValueError: pass
    return None

def norm(s): return re.sub(r'\s+', ' ', str(s)).strip().lower()
def blank(v): return v is None or (isinstance(v, float) and np.isnan(v)) or str(v).strip() == ''
def s(v): return '' if blank(v) else str(v).strip()


# ---------------------------------------------------------------- inputs
def read_settings(path):
    st = pd.read_excel(path, sheet_name='Settings', dtype=object)
    cfg = {str(r.iloc[0]).strip(): ('' if blank(r.iloc[1]) else str(r.iloc[1])) for _, r in st.iterrows()}
    cc = pd.read_excel(path, sheet_name='Cost Centres', dtype=object)
    cfg['_cc'] = {str(a).strip(): str(b).strip() for a, b in zip(cc.iloc[:, 0], cc.iloc[:, 1])}
    rules = pd.read_excel(path, sheet_name='Rules', dtype=object)
    rules.columns = ['prio', 'kw', 'dir', 'exact', 'type', 'acct', 'desc', 'ext', 'cc', 'conf', 'note']
    rules = rules[~rules.kw.map(blank)].copy()
    rules['prio'] = rules.prio.map(num)
    coa = None
    try:
        c = pd.read_excel(path, sheet_name='Chart of Accounts', dtype=object)
        coa = set(c.iloc[:, 0].map(s))
    except Exception:
        pass
    payroll = None
    try:
        raw = pd.read_excel(path, sheet_name='Payroll Schedule', header=None, dtype=object)
        hdr_row = raw.index[raw.iloc[:, 0].astype(str).str.strip() == 'Month (YYYY-MM)'][0]
        pr = raw.iloc[hdr_row + 1:].copy()
        pr.columns = ['month', 'role', 'type', 'amount', 'acct', 'cc', 'desc', 'ext', 'conf']
        pr = pr[pr.acct.map(lambda v: not blank(v))].copy()
        pr['amount_filled'] = ~pr.amount.map(blank)
        pr['amount'] = pr.amount.map(num)
        pr['conf'] = pr.conf.map(lambda v: s(v) or 'High')
        payroll = pr.to_dict('records')
    except Exception:
        pass
    cfg.setdefault('Journal Template Name', 'GENERAL')
    cfg.setdefault('Bank Account Code', 'BNK0001')
    cfg.setdefault('Starting Line No.', '10000')
    return cfg, rules.sort_values('prio').to_dict('records'), coa, payroll

def read_statement(path, sheet):
    raw = pd.read_excel(path, sheet_name=sheet, header=None, dtype=object)
    tx, opening = [], None
    for i, r in raw.iterrows():
        narr = s(r.iloc[C_NARR]) if len(r) > C_NARR else ''
        d = to_date(r.iloc[C_DATE])
        if d is None or not narr: continue
        if 'opening balance' in norm(narr): opening = num(r.iloc[C_WD]) or num(r.iloc[C_DEP]) or None; continue
        if 'closing balance' in norm(narr): continue
        wd, dep = num(r.iloc[C_WD]), num(r.iloc[C_DEP])
        tx.append(dict(row=i + 1, date=d, ref=s(r.iloc[C_REF]), narr=re.sub(r'\s+', ' ', narr), wd=wd, dep=dep, amt=dep - wd))
    return pd.DataFrame(tx), opening


# ---------------------------------------------------------------- classification
def match_rule(narr, amt, rules):
    n = norm(narr)
    for r in rules:
        d = norm(r['dir']) if not blank(r['dir']) else 'any'
        if (d == 'out' and amt > 0) or (d == 'in' and amt < 0): continue
        if not blank(r['exact']) and abs(abs(amt) - num(r['exact'])) > 0.005: continue
        if all(k.strip().lower() in n for k in str(r['kw']).split(';') if k.strip()): return r
    return None

def match_payroll(amt, date, payroll, used):
    # Only Net Salary is matched against the bank statement - it is the one payroll type paid to each
    # employee individually. PAYE/NSSF/SHIF/Housing Levy are bulk on the bank side (see build_payroll_allocation).
    if not payroll or amt >= 0: return None
    month = f'{date.year}-{date.month:02d}'
    cands = [(idx, p) for idx, p in enumerate(payroll)
             if idx not in used and p['type'] == 'SALARY' and p['amount_filled'] and abs(abs(amt) - p['amount']) < 0.005]
    if not cands: return None
    exact_month = [c for c in cands if not blank(c[1].get('month', '')) and s(c[1]['month']) == month]
    idx, p = (exact_month or cands)[0]
    used.add(idx)
    return p

def classify(df, rules, cfg, coa=None, payroll=None):
    prefixes = [p.strip().lower() for p in cfg['Bank charge narrations start with'].split(';') if p.strip()]
    used_payroll = set()
    out = []
    for i, t in df.iterrows():
        rec = dict(t)
        if any(norm(t.narr).startswith(p) for p in prefixes):
            rec.update(kind='charge', status='OK', reason='Bank charge (grouped)'); out.append(rec); continue
        pr = match_payroll(t.amt, t.date, payroll, used_payroll)
        if pr is not None:
            rec.update(kind='txn', acct=s(pr['acct']), desc=s(pr['desc']), ext=s(pr['ext']), cc=s(pr['cc']),
                       status='OK', rule='Payroll Schedule',
                       reason=f"Matched by exact amount to Payroll Schedule row: {s(pr['role'])} {s(pr['type'])}")
            out.append(rec); continue
        r = match_rule(t.narr, t.amt, rules)
        rec.update(kind='txn', acct='', desc='', ext='', cc='', status='UNCLASSIFIED', reason='No matching rule', rule='')
        if r is not None and not blank(r['acct']):
            acct = s(r['acct'])
            if acct == 'SET-BANK-CHARGES': acct = cfg.get('Bank Charges G/L Account No.') or 'SET-ACCOUNT'
            elif acct.startswith('SET-') and acct[4:] in CLEARING_ACCOUNT_SETTING:
                acct = cfg.get(CLEARING_ACCOUNT_SETTING[acct[4:]], '') or 'SET-ACCOUNT'
            rev = (str(r['type']).lower() == 'expense' and t.amt > 0)
            rec.update(acct=acct, desc=s(r['desc']) + (' - Reversal' if rev else ''), ext=s(r['ext']), cc=s(r['cc']),
                       rule=f"P{int(r['prio'])}: {r['kw']}", reason=('Reversal of expense rule; ' if rev else '') + s(r['note']))
            rec['status'] = 'OK' if str(r['conf']).lower() == 'high' else 'CHECK'
            if rec['status'] == 'CHECK': rec['reason'] = f"Rule confidence {r['conf']}. " + rec['reason']
            if coa is not None and acct not in coa and acct not in ('SET-ACCOUNT',):
                rec['status'] = 'CHECK'; rec['reason'] = f"Account {acct} not found in Chart of Accounts sheet - check it. " + rec['reason']
        elif r is not None:
            rec['reason'] = s(r['note']) or 'Review rule'
        out.append(rec)
    res = pd.DataFrame(out)
    # reversal by amount: unclassified inflow equal to earlier outflow(s) that all share one account
    for i, t in res.iterrows():
        if t.kind != 'txn' or t.status != 'UNCLASSIFIED' or t.amt <= 0: continue
        prev = res[(res.index < i) & (res.kind == 'txn') & np.isclose(res.amt, -t.amt)]
        if len(prev) and (prev.status != 'UNCLASSIFIED').all() and prev.acct.nunique() == 1 and not prev.desc.str.contains('Reversal').any():
            p = prev.iloc[-1]
            res.loc[i, ['acct', 'desc', 'ext', 'cc', 'status', 'rule']] = [p.acct, p.desc + ' - Reversal', p.ext, p.cc, 'CHECK', 'Amount match']
            res.loc[i, 'reason'] = f'Inflow equals an earlier payment of {abs(t.amt):,.2f} (statement row {int(p.row)}) - confirm it is a reversal'
    return res


# ---------------------------------------------------------------- journals
def build_journals(res, cfg, start_doc):
    prefix = cfg['Document No. prefix']; width = int(cfg['Document No. digits'] or 6)
    tmpl = cfg['Journal Template Name']; batch = cfg.get('Journal Batch Name') or ''
    bank_acct_code = cfg['Bank Account Code']
    line_no = int(cfg['Starting Line No.'] or 10000)
    rows, meta, n = [], [], start_doc

    def journal(date, ext, acct, desc, cc, amt, status, src_rows, kind='txn'):
        nonlocal n, line_no
        doc = f'{prefix}{n:0{width}d}'; n += 1
        dr_counter = amt < 0
        for side in ('counter', 'bank'):
            is_dr = dr_counter if side == 'counter' else not dr_counter
            acct_type = 'G/L Account' if side == 'counter' else 'Bank Account'
            acct_no = acct if side == 'counter' else bank_acct_code
            rows.append({'Journal Template Name': tmpl, 'Journal Batch Name': batch, 'Line No.': line_no,
                         'Account Type': acct_type, 'Account No.': acct_no, 'Posting Date': date.to_pydatetime(),
                         'Document No.': doc, 'Description': desc,
                         'Debit Amount': round(abs(amt), 2) if is_dr else 0.0,
                         'Credit Amount': 0.0 if is_dr else round(abs(amt), 2),
                         'Shortcut Dimension 1 Code': cc, 'Shortcut Dimension 2 Code': cfg['Project Code'],
                         'Bal. Account Type': acct_type, 'External Document No.': ext, 'Donor': cfg['Donor Code'],
                         '_status': status})
            line_no += 1
        meta.append(dict(doc=doc, rows=src_rows, kind=kind))
        return doc

    docs = {}
    for i, t in res[res.kind == 'txn'].iterrows():
        acct = t.acct or cfg['Suspense G/L Account No.'] or 'REVIEW'
        desc = t.desc or ('REVIEW - ' + t.narr[:60])
        docs[i] = journal(t.date, t.ext or t.ref or 'Unclassified', acct, desc, t.cc, t.amt, t.status, [int(t.row)])

    ch = res[res.kind == 'charge']
    if len(ch):
        mode = cfg['Bank Charges Treatment'].lower(); acct = cfg['Bank Charges G/L Account No.'] or 'SET-ACCOUNT'
        status = 'OK' if cfg['Bank Charges G/L Account No.'] else 'SET ACCOUNT'
        if mode == 'line': groups = [(t.date, [t]) for _, t in ch.iterrows()]
        else:
            key = (lambda d: d) if mode == 'daily' else (lambda d: d.replace(day=calendar.monthrange(d.year, d.month)[1]))
            g = {}
            for _, t in ch.iterrows(): g.setdefault(key(t.date), []).append(t)
            groups = sorted(g.items())
        for date, items in groups:
            net = sum(x.amt for x in items)
            label = 'Bank Charges' + ('' if mode == 'line' else f" - {date:%b %Y}" if mode == 'monthly' else f" - {date:%d %b %Y}")
            journal(date, 'Bank Charges', acct, label, cfg['Bank Charges Cost Centre'], net, status, [int(x.row) for x in items], 'charge')
    gl = pd.DataFrame(rows)
    return gl, pd.DataFrame(meta), docs, n, line_no


# ---------------------------------------------------------------- checks
def build_payroll_allocation(payroll, cfg, start_doc, start_line):
    """One journal per employee per statutory type, straight from the Payroll Schedule - these never touch
    the bank statement (PAYE/NSSF/SHIF/Housing Levy are paid to KRA/NSSF/SHIF as one bulk amount for
    everyone, so there is no individual bank line to match). Only rows where the muster-roll Amount has
    been filled in are generated. Dr the employee's expense account, Cr the type's clearing account."""
    if not payroll: return pd.DataFrame(), 0, start_doc
    prefix = cfg['Document No. prefix']; width = int(cfg['Document No. digits'] or 6)
    tmpl = cfg['Journal Template Name']; batch = cfg.get('Journal Batch Name') or ''
    rows = []; n = start_doc; line_no = start_line
    for p in payroll:
        if p['type'] == 'SALARY' or not p['amount_filled']: continue
        clearing = cfg.get(CLEARING_ACCOUNT_SETTING.get(p['type'], ''), '') or 'SET-ACCOUNT'
        try:
            yr, mo = (int(x) for x in str(p['month']).split('-'))
            date = pd.Timestamp(yr, mo, calendar.monthrange(yr, mo)[1])
        except Exception:
            date = None
        doc = f'{prefix}{n:0{width}d}'; n += 1
        amt = round(abs(p['amount']), 2)
        for side, acct, is_dr in (('expense', s(p['acct']), True), ('clearing', clearing, False)):
            rows.append({'Journal Template Name': tmpl, 'Journal Batch Name': batch, 'Line No.': line_no,
                         'Account Type': 'G/L Account', 'Account No.': acct,
                         'Posting Date': date.to_pydatetime() if date is not None else None,
                         'Document No.': doc, 'Description': s(p['desc']),
                         'Debit Amount': amt if is_dr else 0.0, 'Credit Amount': 0.0 if is_dr else amt,
                         'Shortcut Dimension 1 Code': s(p['cc']), 'Shortcut Dimension 2 Code': cfg['Project Code'],
                         'Bal. Account Type': 'G/L Account', 'External Document No.': s(p['ext']),
                         'Donor': cfg['Donor Code'], '_status': 'OK' if p['conf'] == 'High' else 'CHECK',
                         '_payroll_role': p['role'], '_payroll_type': p['type'],
                         '_no_date': date is None})
            line_no += 1
    return pd.DataFrame(rows), line_no, n


def run_checks(gl, res, cfg, opening):
    chk = []
    by = gl.groupby('Document No.').agg(dr=('Debit Amount', 'sum'), cr=('Credit Amount', 'sum'))
    bad = by[~np.isclose(by.dr, by.cr)]
    chk.append(('Every journal balances (debits = credits)', 'PASS' if bad.empty else 'FAIL', f'{len(bad)} unbalanced'))
    b = gl[gl['Account Type'] == 'Bank Account']
    gl_net = b['Debit Amount'].sum() - b['Credit Amount'].sum(); st_net = res.amt.sum()
    chk.append(('Bank movement in GL = statement movement', 'PASS' if abs(gl_net - st_net) < 0.005 else 'FAIL',
                f'Statement {st_net:,.2f} | GL {gl_net:,.2f} | difference {gl_net - st_net:,.2f}'))
    dup = res[res.duplicated(['date', 'ref', 'narr', 'amt'], keep=False)]
    chk.append(('No duplicate statement lines', 'PASS' if dup.empty else 'WARN', f'{len(dup)} suspicious lines'))
    un = int((res.status == 'UNCLASSIFIED').sum())
    chk.append(('All lines coded to a real account', 'PASS' if un == 0 else 'ACTION', f'{un} unclassified lines need an account'))
    chk.append(('Bank charges account set', 'PASS' if cfg['Bank Charges G/L Account No.'] else 'ACTION', 'Set it in GL_Rules.xlsx > Settings'))
    chk.append(('Suspense account set', 'PASS' if cfg['Suspense G/L Account No.'] else 'ACTION', 'Set it in GL_Rules.xlsx > Settings'))
    chk.append(('Opening balance on statement', 'INFO', 'Not provided - closing balance cannot be proven' if opening is None else f'{opening:,.2f}'))
    return chk


# ---------------------------------------------------------------- test against posted GL
def read_posted(path, sheet, cfg):
    g = pd.read_excel(path, sheet_name=sheet)
    g['Posting Date'] = g['Posting Date'].apply(lambda v: to_date(v) if isinstance(v, str) else pd.Timestamp(v))
    for c in ('Debit Amount', 'Credit Amount'): g[c] = g[c].fillna(0)
    bank = cfg['Bank G/L Account No.']; J = []
    for doc, d in g.groupby('Document No.', sort=False):
        isbank = (d['G/L Account No.'] == bank) | (d['G/L Account No.'].isna() & d['Bal. Account Type'].astype(str).str.startswith('Bank'))
        bk, ex = d[isbank], d[~isbank]
        e = ex.iloc[0] if len(ex) else d.iloc[0]
        J.append(dict(doc=doc, date=d['Posting Date'].iloc[0], bank=(bk['Debit Amount'] - bk['Credit Amount']).sum(),
                      dr=d['Debit Amount'].sum(), cr=d['Credit Amount'].sum(),
                      acct=s(e['G/L Account No.']), desc=s(e['Description']), ext=s(e['External Document No.']), cc=s(e['cost Center Code']),
                      proj=s(d['Project Code'].iloc[0]), donor=s(d['Donor code '].iloc[0]),
                      bank_acct_missing=bool(bk['G/L Account No.'].isna().any()), ex_acct_missing=bool(len(ex) and ex['G/L Account No.'].isna().any()),
                      types=set(d['Bal. Account Type'].astype(str).str.strip())))
    return pd.DataFrame(J), g

def compare(res, posted, tool_docs):
    tx = res[res.kind == 'txn']; used = set(); pairs = []
    for i, t in tx.iterrows():
        m = posted[(~posted.doc.isin(used)) & np.isclose(posted.bank, t.amt)]
        if m.empty: pairs.append((i, None)); continue
        m = m.assign(dd=(m.date - t.date).abs()).sort_values('dd'); p = m.iloc[0]
        if p.dd.days > 3: pairs.append((i, None)); continue
        used.add(p.doc); pairs.append((i, p))
    lines = []
    for i, p in pairs:
        t = res.loc[i]
        rec = dict(row=int(t.row), date=t.date, narr=t.narr, amt=t.amt, tool_status=t.status, tool_acct=t.acct, tool_desc=t.desc, tool_cc=t.cc, tool_ext=t.ext)
        if p is None: rec.update(posted_doc='', posted_acct='', posted_desc='', posted_cc='', posted_ext='')
        else: rec.update(posted_doc=p.doc, posted_acct=p.acct, posted_desc=p.desc, posted_cc=p.cc, posted_ext=p.ext)
        lines.append(rec)
    L = pd.DataFrame(lines)
    L['comparable'] = (L.posted_doc != '') & (L.posted_acct != '')
    L['acct_ok'] = L.comparable & (L.tool_acct == L.posted_acct)
    L['cc_ok'] = L.comparable & (L.tool_cc == L.posted_cc)
    L['desc_ok'] = L.comparable & (L.tool_desc.str.strip().str.lower() == L.posted_desc.str.strip().str.lower())
    L['ext_ok'] = L.comparable & (L.tool_ext.str.strip().str.lower() == L.posted_ext.str.strip().str.lower())
    L['verdict'] = np.where(L.posted_doc == '', 'Not booked in posted GL',
                    np.where(L.posted_acct == '', 'Posted entry has no account',
                    np.where(L.tool_status == 'UNCLASSIFIED', 'Tool sent to review',
                    np.where(L.acct_ok, 'Account matches', 'Account DIFFERS'))))
    return L, used


def findings(posted, raw_posted, res, used, cfg):
    F = []
    for _, j in posted.iterrows():
        if not np.isclose(j.dr, j.cr): F.append((j.doc, 'Journal does not balance', f'Debits {j.dr:,.2f} vs credits {j.cr:,.2f} (difference {j.dr - j.cr:,.2f})'))
        if j.ex_acct_missing or j.bank_acct_missing: F.append((j.doc, 'Line has no G/L account', f'{j.desc}: account blank on posted line'))
        if not j.proj or not j.donor: F.append((j.doc, 'Project / donor code missing', j.desc))
    bank = cfg['Bank G/L Account No.']
    bad_types = raw_posted[raw_posted['G/L Account No.'].notna() & (raw_posted['G/L Account No.'] != bank) & raw_posted['Bal. Account Type'].astype(str).str.startswith('Bank')]
    for d in bad_types['Document No.'].unique(): F.append((d, 'Counter line marked "Bank Account"', 'Bal. Account Type is inconsistent with the other journals'))
    return pd.DataFrame(F, columns=['Document', 'Finding', 'Detail'])


# ---------------------------------------------------------------- writers
def write_df(ws, df, widths=None, start_row=1):
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(start_row, j, c); cell.fill = NAVY; cell.font = WHITE_B; cell.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (pd.Timestamp,)): v = v.to_pydatetime()
            if isinstance(v, (np.bool_,)): v = bool(v)
            if isinstance(v, (np.integer,)): v = int(v)
            if isinstance(v, (np.floating,)): v = None if np.isnan(v) else float(v)
            c = ws.cell(i, j, v)
            if isinstance(v, dt.datetime): c.number_format = 'DD-MM-YYYY'
            elif isinstance(v, float): c.number_format = '#,##0.00'
    for j, w in enumerate(widths or [], 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)


def write_gl(gl, path):
    from bc_template import BC_ROW1, BC_HEADERS
    wb = Workbook(); ws = wb.active; ws.title = 'Gen. Journal Line'
    for j, v in enumerate(BC_ROW1, 1): ws.cell(1, j, v)
    for j, v in enumerate(BC_HEADERS, 1):
        c = ws.cell(3, j, v); c.fill = NAVY; c.font = WHITE_B
    col_of = {h: j + 1 for j, h in enumerate(BC_HEADERS)}
    date_col = col_of['Posting Date']; dr_col = col_of['Debit Amount']; cr_col = col_of['Credit Amount']
    max_used_col = max(col_of[h] for h in ['Journal Template Name', 'Journal Batch Name', 'Line No.', 'Account Type',
                        'Account No.', 'Posting Date', 'Document No.', 'Description', 'Debit Amount', 'Credit Amount',
                        'Shortcut Dimension 1 Code', 'Shortcut Dimension 2 Code', 'Bal. Account Type',
                        'External Document No.', 'Donor'])
    r = 4
    for _, rec in gl.iterrows():
        for h, j in col_of.items():
            if h in rec and not (isinstance(rec[h], float) and np.isnan(rec[h]) if h not in ('Debit Amount', 'Credit Amount') else False):
                ws.cell(r, j, rec[h])
        ws.cell(r, date_col).number_format = 'DD-MM-YYYY'
        ws.cell(r, dr_col).number_format = '#,##0.00'; ws.cell(r, cr_col).number_format = '#,##0.00'
        fill = {'CHECK': YELLOW, 'UNCLASSIFIED': ORANGE, 'SET ACCOUNT': ORANGE}.get(rec['_status'])
        if fill:
            for c in range(1, max_used_col + 1): ws.cell(r, c).fill = fill
        r += 1
    widths = {'Journal Template Name': 12, 'Journal Batch Name': 14, 'Line No.': 9, 'Account Type': 13,
              'Account No.': 13, 'Posting Date': 13, 'Document No.': 12, 'Description': 44, 'Debit Amount': 15,
              'Credit Amount': 15, 'Shortcut Dimension 1 Code': 13, 'Shortcut Dimension 2 Code': 13,
              'Bal. Account Type': 15, 'External Document No.': 32, 'Donor': 10}
    for h, w in widths.items(): ws.column_dimensions[get_column_letter(col_of[h])].width = w
    ws.freeze_panes = 'A4'
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--statement', required=True); ap.add_argument('--sheet', required=True)
    ap.add_argument('--rules', required=True); ap.add_argument('--start-doc', type=int, required=True, help='first document number, e.g. 386 for JV000386')
    ap.add_argument('--compare-gl', default=None, help='sheet name of already-posted GL in the same workbook (for testing)')
    ap.add_argument('--out-dir', default='.'); ap.add_argument('--tag', default='')
    a = ap.parse_args()
    out = Path(a.out_dir); out.mkdir(parents=True, exist_ok=True)

    cfg, rules, coa, payroll = read_settings(a.rules)
    stmt, opening = read_statement(a.statement, a.sheet)
    res = classify(stmt, rules, cfg, coa, payroll)
    gl, meta, docs, next_doc, next_line = build_journals(res, cfg, a.start_doc)
    res['doc'] = [docs.get(i, '') for i in res.index]
    pay_gl, _, _ = build_payroll_allocation(payroll, cfg, next_doc, next_line)
    if len(pay_gl): gl = pd.concat([gl, pay_gl], ignore_index=True)
    checks = run_checks(gl, res, cfg, opening)
    tag = a.tag or Path(a.sheet).name.replace(' ', '_')
    write_gl(gl, out / f'GL_Import_{tag}_DRAFT.xlsx')

    # ---- review workbook
    wb = Workbook(); ws = wb.active; ws.title = 'Summary'
    tx = res[res.kind == 'txn']
    S = [('BANK STATEMENT -> GL: RUN SUMMARY', ''), ('Statement lines read', len(res)),
         ('  Bank charge lines (grouped into journal(s))', int((res.kind == 'charge').sum())), ('  Transaction lines', len(tx)),
         ('  Coded automatically, high confidence', int((tx.status == 'OK').sum())),
         ('  Coded but please check (medium/low confidence, reversal guesses)', int((tx.status == 'CHECK').sum())),
         ('  Unclassified (need your input)', int((tx.status == 'UNCLASSIFIED').sum())),
         ('Journals created', gl['Document No.'].nunique()), ('Total money out (statement)', float(res.wd.sum())), ('Total money in (statement)', float(res.dep.sum())),
         ('', ''), ('Payroll Allocation (PAYE/NSSF/SHIF/Housing Levy - not from the bank statement)', ''),
         ('  Generated this run (muster roll amount was filled in)', int(len(pay_gl) // 2)),
         ('  Still blank in Payroll Schedule (muster roll not filled in yet)', int(sum(1 for p in (payroll or []) if p['type'] != 'SALARY' and not p['amount_filled']))), ('', '')]
    for k, v in S: ws.append([k, v])
    ws.append(['CHECKS', 'Result', 'Detail']);
    for c in ws[ws.max_row]: c.fill = NAVY; c.font = WHITE_B
    for c in checks:
        ws.append(list(c)); cell = ws.cell(ws.max_row, 2)
        cell.fill = {'PASS': GREEN, 'FAIL': ORANGE, 'ACTION': YELLOW, 'WARN': YELLOW}.get(c[1], PatternFill())
    ws['A1'].font = Font(bold=True, size=14)
    for r in range(2, 12): ws.cell(r, 2).number_format = '#,##0.00' if r in (9, 10) else '0'
    ws.column_dimensions['A'].width = 66; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 80

    rv = res[res.status.isin(['CHECK', 'UNCLASSIFIED'])].copy()
    rv['Action'] = np.where(rv.status == 'UNCLASSIFIED', 'Provide account (add a rule after)', 'Confirm or correct')
    rv = rv[['doc', 'date', 'ref', 'narr', 'amt', 'status', 'acct', 'desc', 'cc', 'reason', 'Action']]
    rv.columns = ['Document No.', 'Date', 'Reference', 'Narration', 'Amount (+in / -out)', 'Status', 'Proposed Account', 'Proposed Description', 'Cost Centre', 'Why flagged', 'Action']
    write_df(wb.create_sheet('Needs Review'), rv, [12, 12, 20, 62, 16, 14, 14, 40, 11, 70, 30])
    al = res[['doc', 'row', 'date', 'ref', 'narr', 'wd', 'dep', 'kind', 'status', 'acct', 'desc', 'ext', 'cc', 'rule', 'reason']].copy()
    al.columns = ['Document No.', 'Statement Row', 'Date', 'Reference', 'Narration', 'Money Out', 'Money In', 'Type', 'Status', 'Account', 'Description', 'External Doc', 'Cost Centre', 'Rule used', 'Note']
    write_df(wb.create_sheet('All Lines'), al, [12, 9, 12, 20, 62, 13, 13, 8, 14, 12, 40, 40, 11, 26, 50])

    if len(pay_gl):
        pa = pay_gl[pay_gl['Account Type'] == 'G/L Account'].copy()
        pa = pa[pa['Debit Amount'] > 0][['Document No.', '_payroll_role', '_payroll_type', 'Posting Date', 'Debit Amount',
                                          'Account No.', 'Shortcut Dimension 1 Code', '_status', '_no_date']]
        pa.columns = ['Document No.', 'Employee', 'Type', 'Posting Date', 'Amount', 'Account', 'Cost Centre', 'Status', 'No date in schedule']
        write_df(wb.create_sheet('Payroll Allocation'), pa, [12, 30, 10, 13, 14, 14, 12, 10, 16])
    blank_pay = [(p['role'], p['type']) for p in (payroll or []) if p['type'] != 'SALARY' and not p['amount_filled']]
    if blank_pay:
        wp = wb.create_sheet('Payroll Not Yet Filled')
        wp.append(['Role', 'Type', 'Note'])
        for c in wp[1]: c.fill = NAVY; c.font = WHITE_B
        for role, typ in blank_pay: wp.append([role, typ, 'Fill in the Amount column in Payroll Schedule, then re-run'])
        wp.column_dimensions['A'].width = 32; wp.column_dimensions['B'].width = 10; wp.column_dimensions['C'].width = 45

    # ---- test vs posted GL
    if a.compare_gl:
        posted, raw = read_posted(a.statement, a.compare_gl, cfg)
        L, used = compare(res, posted, docs)
        cmpb = L[L.comparable]; auto = cmpb[cmpb.tool_status != 'UNCLASSIFIED']; high = auto[auto.tool_status == 'OK']
        pct = lambda x, n: f'{x}/{n} = {x / n:.0%}' if n else 'n/a'
        T = [('TEST: TOOL vs YOUR POSTED ENTRIES', ''),
             ('IMPORTANT', 'Rules were learned from this same month, so these results are optimistic. The honest test is a NEW month (e.g. February).'), ('', ''),
             ('Transaction lines on the statement (excl. bank charges)', len(L)),
             ('  found in your posted GL (matched on amount/date)', int((L.posted_doc != '').sum())),
             ('  with a usable posted account to compare against', len(cmpb)),
             ('Tool coded automatically (any confidence)', pct(len(auto), len(cmpb))),
             ('  account agrees with your posting', pct(int(auto.acct_ok.sum()), len(auto))),
             ('  cost centre agrees', pct(int(auto.cc_ok.sum()), len(auto))),
             ('  description agrees exactly', pct(int(auto.desc_ok.sum()), len(auto))),
             ('  external doc no. agrees', pct(int(auto.ext_ok.sum()), len(auto))),
             ('High-confidence lines only: account agrees', pct(int(high.acct_ok.sum()), len(high))),
             ('Tool sent to review, but you posted an account', int(((L.verdict == 'Tool sent to review')).sum()))]
        wt = wb.create_sheet('Test vs Posted')
        for k, v in T: wt.append([k, v])
        wt['A1'].font = Font(bold=True, size=14); wt['A2'].font = Font(bold=True, color='C00000')
        wt.column_dimensions['A'].width = 62; wt.column_dimensions['B'].width = 95
        r0 = wt.max_row + 2
        Lo = L[['row', 'date', 'narr', 'amt', 'verdict', 'tool_status', 'tool_acct', 'posted_acct', 'tool_desc', 'posted_desc', 'tool_cc', 'posted_cc', 'posted_doc']].copy()
        Lo.columns = ['Stmt Row', 'Date', 'Narration', 'Amount', 'Verdict', 'Tool Status', 'Tool Account', 'Your Account', 'Tool Description', 'Your Description', 'Tool CC', 'Your CC', 'Your Doc']
        write_df(wt, Lo, None, start_row=r0)
        for j, w in enumerate([62, 95, 60, 14, 26, 13, 14, 14, 40, 40, 10, 10, 11], 1): wt.column_dimensions[get_column_letter(j)].width = max(w, 12) if j > 2 else wt.column_dimensions[get_column_letter(j)].width
        wt.freeze_panes = None
        for r in range(r0 + 1, wt.max_row + 1):
            v = wt.cell(r, 5).value
            fill = GREEN if v == 'Account matches' else ORANGE if v in ('Account DIFFERS',) else YELLOW
            wt.cell(r, 5).fill = fill

        # bridge: statement vs posted GL
        st_only = tx[~tx.index.isin(L[L.posted_doc != ''].index if False else [])]
        matched_rows = set(L[L.posted_doc != '']['row']); st_only = tx[~tx.row.isin(matched_rows)]
        gl_only = posted[~posted.doc.isin(used)]
        chg = res[res.kind == 'charge']
        B = []
        for _, t in st_only.iterrows(): B.append(('On statement, not in posted GL', f"Stmt row {int(t.row)}", t.date, t.narr[:70], t.amt))
        for _, j in gl_only.iterrows(): B.append(('In posted GL, no exact statement match', j.doc, j.date, j.desc, -j.bank))
        B.append(('Bank charges not in posted GL', f'{len(chg)} lines', None, 'Charges, excise duty, SMS fees, net of refunds', chg.amt.sum()))
        def likely(kind, ref, date, amt):
            # pair a GL-only journal with a statement-only line on the same day with a near / 10x amount (typo finder)
            if kind.startswith('In posted GL'):
                c = st_only[((st_only.date - date).abs() <= pd.Timedelta(days=1))]
                for _, t2 in c.iterrows():
                    r = abs(t2.amt) / abs(amt) if amt else 0
                    if abs(r - 1) <= 0.01 or abs(r - 10) < 1e-6 or abs(r - 0.1) < 1e-6:
                        return f"Likely typo: statement row {int(t2.row)} is {abs(t2.amt):,.2f}, GL has {abs(amt):,.2f}"
            return ''
        Bd = pd.DataFrame(B, columns=['Reconciling item', 'Ref', 'Date', 'Detail', 'Effect on (statement - GL) movement'])
        Bd['Likely explanation'] = [likely(r[0], r[1], r[2], (-r[4] if r[0].startswith('In posted GL') else r[4])) for r in B]
        st_net = res.amt.sum(); gl_net = posted.bank.sum()
        wb2 = wb.create_sheet('Bank Rec Bridge')
        wb2.append(['Statement net movement', st_net]); wb2.append(['Posted GL net bank movement', gl_net]); wb2.append(['Difference to explain', st_net - gl_net])
        wb2.append(['Sum of reconciling items below', float(Bd.iloc[:, 4].sum())]); wb2.append(['Unexplained', round(st_net - gl_net - float(Bd.iloc[:, 4].sum()), 2)])
        for r in range(1, 6): wb2.cell(r, 2).number_format = '#,##0.00'; wb2.cell(r, 1).font = Font(bold=True)
        write_df(wb2, Bd, [46, 16, 13, 72, 34, 70], start_row=7); wb2.freeze_panes = None
        F = findings(posted, raw, res, used, cfg)
        write_df(wb.create_sheet('Findings in Posted GL'), F, [12, 42, 90])

    wb.save(out / f'Review_{tag}.xlsx')
    print('Journals:', gl['Document No.'].nunique(), '| checks:', [(c[0][:28], c[1]) for c in checks])


if __name__ == '__main__':
    main()
