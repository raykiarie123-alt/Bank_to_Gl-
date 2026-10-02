"""Bank statement -> GL ledger entries. Simple upload screen (no AI, runs on your own PC)."""
import subprocess, sys, tempfile
from pathlib import Path
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

HERE = Path(__file__).parent
ENGINE = HERE / "bank_to_gl.py"
DEFAULT_RULES = HERE / "GL_Rules.xlsx"
OUT = HERE / "outputs"; OUT.mkdir(exist_ok=True)

st.set_page_config(page_title="Bank Statement to Ledger", page_icon="📒", layout="wide")
st.title("📒 Bank Statement → Ledger Entries")
st.caption("Upload the statement, click the button, review, then download. Nothing is posted to D365 automatically.")

st.subheader("Step 1. Upload the bank statement")
up = st.file_uploader("Excel (.xlsx) or CSV", type=["xlsx", "xlsm", "csv"])

tmp = Path(tempfile.mkdtemp())
stmt_path, sheet = None, None
if up is not None:
    if up.name.lower().endswith(".csv"):
        stmt_path = tmp / "statement.xlsx"
        pd.read_csv(up, header=None, dtype=str).to_excel(stmt_path, header=False, index=False, sheet_name="Statement")
        sheet = "Statement"
    else:
        stmt_path = tmp / "statement.xlsx"; stmt_path.write_bytes(up.getbuffer())
        names = load_workbook(stmt_path, read_only=True).sheetnames
        sheet = st.selectbox("Which sheet holds the bank statement?", names)

st.subheader("Step 2. Details")
c1, c2 = st.columns(2)
start = c1.number_input("Next journal number to use (e.g. 386 for JV000386)", min_value=1, value=1, step=1)
tag = c2.text_input("Name for the output files", value="Statement")
rules_up = st.file_uploader("Optional: use a different rules file (otherwise GL_Rules.xlsx in this folder is used)", type=["xlsx"], key="rules")
rules_path = DEFAULT_RULES
if rules_up is not None:
    rules_path = tmp / "rules.xlsx"; rules_path.write_bytes(rules_up.getbuffer())

st.subheader("Step 3. Prepare")
if st.button("Prepare ledger entries", type="primary", disabled=stmt_path is None):
    cmd = [sys.executable, str(ENGINE), "--statement", str(stmt_path), "--sheet", sheet, "--rules", str(rules_path),
           "--start-doc", str(int(start)), "--out-dir", str(OUT), "--tag", tag.replace(" ", "_")]
    p = subprocess.run(cmd, capture_output=True, text=True)
    if p.returncode != 0:
        st.error("Something went wrong. Check that the statement layout is the same as usual "
                 "(date in column A, reference in C, narration in D, money out in F, money in in G).")
        st.code(p.stderr[-1500:])
    else:
        t = tag.replace(" ", "_")
        gl_file, rv_file = OUT / f"GL_Import_{t}_DRAFT.xlsx", OUT / f"Review_{t}.xlsx"
        summ = pd.read_excel(rv_file, sheet_name="Summary", header=None).fillna("")
        st.success("Done. Please review before importing.")
        st.dataframe(summ, hide_index=True, use_container_width=True)
        st.subheader("Lines that need your attention")
        st.dataframe(pd.read_excel(rv_file, sheet_name="Needs Review"), hide_index=True, use_container_width=True)
        d1, d2 = st.columns(2)
        d1.download_button("⬇ Download ledger entries (import file)", gl_file.read_bytes(), gl_file.name)
        d2.download_button("⬇ Download review report", rv_file.read_bytes(), rv_file.name)
        st.info("Yellow/orange rows in the import file need checking. Fix them in Excel, then import into D365 as an UNPOSTED journal.")
